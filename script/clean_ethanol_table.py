import re
from pathlib import Path
import json

import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


SCRIPT_DIR = Path(__file__).resolve().parent
PROJECT_ROOT = SCRIPT_DIR.parent
CONFIG_PATH = PROJECT_ROOT / "private_inputs.json"
LOCAL_CONFIG = json.loads(CONFIG_PATH.read_text()) if CONFIG_PATH.exists() else {}
RAW_DIR = PROJECT_ROOT / "raw_data"
OUTPUT_DIR = PROJECT_ROOT / "clean_data"

INPUT_FILE = RAW_DIR / LOCAL_CONFIG.get("ethanol_file", "ethanol.xlsx")
OUT_XLSX = OUTPUT_DIR / "ASPREE_ethanol_contamination_clean.xlsx"
OUT_CSV = OUTPUT_DIR / "ASPREE_ethanol_contamination_clean.csv"


def slugify(value):
    value = str(value).strip().lower()
    value = value.replace("%", "pct")
    value = re.sub(r"[^a-z0-9]+", "_", value)
    return value.strip("_")


def build_column_name(raw1, raw2):
    raw1 = raw1 or ""
    raw2 = raw2 or ""

    if raw2 == "LS Accession":
        return "labcorp_accession"
    if raw2 == "Client Accession":
        return "subject_id"

    if raw1:
        return f"{slugify(raw1)}__{slugify(raw2)}"
    return slugify(raw2)


def fill_header_row(values):
    filled = []
    current = None
    for value in values:
        if value not in (None, ""):
            current = value
            filled.append(value)
        else:
            filled.append(current)
    return filled


def auto_width(ws, max_width=28):
    for col_idx in range(1, ws.max_column + 1):
        letter = get_column_letter(col_idx)
        max_len = max(
            (len(str(cell.value)) for cell in ws[get_column_letter(col_idx)] if cell.value is not None),
            default=8,
        )
        ws.column_dimensions[letter].width = min(max_len + 2, max_width)


def style_sheet(ws):
    fill = PatternFill("solid", fgColor="D9EAF7")
    thin = Side(style="thin", color="B7C9D6")
    for cell in ws[1]:
        cell.fill = fill
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=ws.max_column):
        for cell in row:
            cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"
    auto_width(ws)


def main():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    ws = openpyxl.load_workbook(INPUT_FILE, data_only=True).active

    raw1 = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
    raw1 = fill_header_row(raw1)
    raw2 = [ws.cell(2, c).value for c in range(1, ws.max_column + 1)]
    headers = [build_column_name(raw1[c - 1], raw2[c - 1]) for c in range(1, ws.max_column + 1)]

    data_rows = []
    for r in range(3, ws.max_row + 1):
        row = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
        if any(value is not None for value in row):
            data_rows.append(row)

    wb_out = Workbook()
    ws_out = wb_out.active
    ws_out.title = "Clean"
    ws_out.append(headers)
    for row in data_rows:
        ws_out.append(row)
    style_sheet(ws_out)
    wb_out.save(OUT_XLSX)

    import csv

    with OUT_CSV.open("w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(headers)
        writer.writerows(data_rows)

    print("Rows:", len(data_rows))
    print("Columns:", len(headers))
    print("Outputs:")
    print(OUT_XLSX.relative_to(PROJECT_ROOT))
    print(OUT_CSV.relative_to(PROJECT_ROOT))
    print("Headers:")
    for header in headers:
        print(header)


if __name__ == "__main__":
    main()
