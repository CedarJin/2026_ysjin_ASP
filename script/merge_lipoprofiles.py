import math
from pathlib import Path
import json

import openpyxl
import xlrd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


SCRIPT_DIR = Path(__file__).resolve().parent
PROJECT_ROOT = SCRIPT_DIR.parent
CONFIG_PATH = PROJECT_ROOT / "private_inputs.json"
LOCAL_CONFIG = json.loads(CONFIG_PATH.read_text()) if CONFIG_PATH.exists() else {}
RAW_DIR = PROJECT_ROOT / "raw_data"
OUTPUT_DIR = PROJECT_ROOT / "clean_data"

NMR_FILE = RAW_DIR / LOCAL_CONFIG.get("nmr_file", "nmr_lipoprofile.xls")
LIPID_FILE = RAW_DIR / LOCAL_CONFIG.get("lipid_file", "lipids_report.xlsx")

OUT_NICE = OUTPUT_DIR / "ASPREE_merged_lipoprofile_stratified.xlsx"
OUT_CLEAN_XLSX = OUTPUT_DIR / "ASPREE_merged_lipoprofile_clean.xlsx"
OUT_CLEAN_CSV = OUTPUT_DIR / "ASPREE_merged_lipoprofile_clean.csv"
OUT_CODEBOOK = OUTPUT_DIR / "ASPREE_merged_lipoprofile_codebook.xlsx"


def normalize_header_value(value):
    if value is None:
        return ""
    if isinstance(value, float) and math.isnan(value):
        return ""
    return str(value).strip()


def normalize_data_value(value):
    if value is None:
        return None
    if isinstance(value, float) and math.isnan(value):
        return None
    if isinstance(value, str):
        value = value.strip()
        return value if value != "" else None
    return value


def read_xls(path):
    sheet = xlrd.open_workbook(path).sheet_by_index(0)
    rows = []
    for r in range(sheet.nrows):
        rows.append([sheet.cell_value(r, c) for c in range(sheet.ncols)])
    return rows


def read_xlsx(path):
    ws = openpyxl.load_workbook(path, data_only=True).active
    rows = []
    for r in ws.iter_rows(values_only=True):
        rows.append(list(r))
    return rows


def build_columns(rows):
    header_rows = rows[:5]
    data_rows = rows[5:]
    columns = []
    for idx in range(len(header_rows[0])):
        col = {
            "study": normalize_header_value(header_rows[0][idx] if idx < len(header_rows[0]) else ""),
            "group": normalize_header_value(header_rows[1][idx] if idx < len(header_rows[1]) else ""),
            "name": normalize_header_value(header_rows[2][idx] if idx < len(header_rows[2]) else ""),
            "unit": normalize_header_value(header_rows[3][idx] if idx < len(header_rows[3]) else ""),
            "abbr": normalize_header_value(header_rows[4][idx] if idx < len(header_rows[4]) else ""),
            "values": [normalize_data_value(row[idx] if idx < len(row) else None) for row in data_rows],
        }
        columns.append(col)
    return columns


def enrich_for_codebook(columns):
    study = ""
    group = ""
    enriched = []
    for idx, col in enumerate(columns, start=1):
        if col["study"]:
            study = col["study"]
        if col["group"]:
            group = col["group"]

        variable_name = col["name"] or col["abbr"]
        variable_group = group
        study_name = study

        if idx == 1:
            variable_name = "Labcorp Accession Number"
            variable_group = "Identifiers"
        elif idx == 2:
            variable_name = "Subject ID"
            variable_group = "Identifiers"
        elif idx == 3:
            variable_name = "Collection Date"
            variable_group = "Identifiers"

        unit = col["unit"]
        if idx <= 3 or unit == "Measurement Units =>":
            unit = ""

        enriched.append(
            {
                "abbr": col["abbr"],
                "name": variable_name,
                "group": variable_group,
                "unit": unit,
                "study": study_name,
            }
        )
    return enriched


def same_series(values1, values2):
    if len(values1) != len(values2):
        return False
    for left, right in zip(values1, values2):
        if left is None and right is None:
            continue
        if left is None or right is None:
            return False
        try:
            if float(left) != float(right):
                return False
        except Exception:
            if str(left) != str(right):
                return False
    return True


def write_clean_csv(path, headers, data_rows):
    import csv

    with path.open("w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(headers)
        writer.writerows(data_rows)


def auto_width(ws, max_width=28):
    for col_idx in range(1, ws.max_column + 1):
        letter = get_column_letter(col_idx)
        values = []
        for row in ws.iter_rows(min_col=col_idx, max_col=col_idx, values_only=True):
            value = row[0]
            if value is None:
                continue
            values.append(len(str(value)))
        width = min(max(values, default=8) + 2, max_width)
        ws.column_dimensions[letter].width = width


def merge_repeated_cells(ws, row_idx, start_col, end_col):
    current_start = start_col
    current_value = ws.cell(row_idx, start_col).value
    for col in range(start_col + 1, end_col + 2):
        value = ws.cell(row_idx, col).value if col <= end_col else None
        if value != current_value:
            if current_value not in (None, "") and col - 1 > current_start:
                ws.merge_cells(
                    start_row=row_idx,
                    start_column=current_start,
                    end_row=row_idx,
                    end_column=col - 1,
                )
            current_start = col
            current_value = value


def style_stratified_sheet(ws, total_rows, total_cols):
    fill_title = PatternFill("solid", fgColor="1F4E78")
    fill_header = PatternFill("solid", fgColor="D9EAF7")
    fill_meta = PatternFill("solid", fgColor="EAF4E2")
    thin = Side(style="thin", color="B7C9D6")

    for cell in ws[1]:
        cell.fill = fill_title
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row_idx in (2, 3, 4, 5):
        for cell in ws[row_idx]:
            cell.fill = fill_header
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for row in range(6, total_rows + 1):
        for col in range(1, min(4, total_cols + 1)):
            ws.cell(row, col).fill = fill_meta

    for row in ws.iter_rows(min_row=1, max_row=total_rows, min_col=1, max_col=total_cols):
        for cell in row:
            cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
            if cell.row >= 6 and cell.column >= 4:
                cell.alignment = Alignment(horizontal="center")

    ws.freeze_panes = "D6"
    ws.auto_filter.ref = f"A5:{get_column_letter(total_cols)}{total_rows}"
    auto_width(ws, max_width=24)


def is_comment_column(col):
    candidates = [col.get("abbr", ""), col.get("name", ""), col.get("group", "")]
    return any("comment" in str(value).strip().lower() for value in candidates if value is not None)


def reorder_comments_last(columns):
    non_comments = [col for col in columns if not is_comment_column(col)]
    comments = [col for col in columns if is_comment_column(col)]
    return non_comments + comments


def move_columns_to_end(columns, abbreviations):
    targets = [col for col in columns if col.get("abbr") in abbreviations]
    remaining = [col for col in columns if col.get("abbr") not in abbreviations]
    ordered_targets = []
    for abbr in abbreviations:
        ordered_targets.extend([col for col in targets if col.get("abbr") == abbr])
    return remaining + ordered_targets


def apply_header_overrides(columns):
    for col in columns:
        abbr = col.get("abbr")
        if abbr == "ApoA1":
            col["group"] = "Derived Apolipoprotein Concentrations"
            col["name"] = "ApoA-1"
            col["unit"] = "mg/dL"
        elif abbr == "NLDLC":
            col["name"] = "non-LDL cholesterol"
        elif abbr == "NHDLC":
            col["name"] = "non-HDL cholesterol"
        elif abbr == "ApoB":
            col["group"] = "Derived Apolipoprotein Concentrations"
    return columns


def merge_header_row_with_fill(ws, row_idx, start_col, end_col):
    current_label = None
    run_start = None
    for col in range(start_col, end_col + 1):
        raw_value = ws.cell(row_idx, col).value
        value = normalize_header_value(raw_value)
        if value:
            if current_label == value:
                continue
            if current_label is not None and col - 1 > run_start:
                ws.merge_cells(
                    start_row=row_idx,
                    start_column=run_start,
                    end_row=row_idx,
                    end_column=col - 1,
                )
            current_label = value
            run_start = col
        elif current_label is not None:
            continue
        else:
            run_start = None

    if current_label is not None and run_start is not None and end_col > run_start:
        ws.merge_cells(
            start_row=row_idx,
            start_column=run_start,
            end_row=row_idx,
            end_column=end_col,
        )


def style_simple_sheet(ws):
    fill = PatternFill("solid", fgColor="D9EAF7")
    thin = Side(style="thin", color="B7C9D6")
    for cell in ws[1]:
        cell.fill = fill
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=ws.max_column):
        for cell in row:
            cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"
    auto_width(ws)


def main():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    nmr_rows = read_xls(NMR_FILE)
    lipid_rows = read_xlsx(LIPID_FILE)

    nmr_cols = build_columns(nmr_rows)
    lipid_cols = build_columns(lipid_rows)

    # Remove AL:AR from the NMR workbook, i.e. 1-based AL(38) to AR(44).
    keep_nmr = [col for idx, col in enumerate(nmr_cols, start=1) if not (38 <= idx <= 44)]

    nmr_abbrs = {col["abbr"]: col for col in keep_nmr if col["abbr"]}
    add_cols = []
    duplicate_report = []
    for col in lipid_cols[3:]:
        abbr = col["abbr"]
        if abbr in nmr_abbrs:
            duplicate_report.append(abbr)
            if same_series(nmr_abbrs[abbr]["values"], col["values"]):
                continue
        add_cols.append(col)

    merged_cols = keep_nmr + add_cols
    merged_cols = move_columns_to_end(merged_cols, ["ApoA1", "ApoB"])
    merged_cols = reorder_comments_last(merged_cols)
    merged_cols = apply_header_overrides(merged_cols)
    data_row_count = len(merged_cols[0]["values"])

    wb_nice = Workbook()
    ws_nice = wb_nice.active
    ws_nice.title = "Merged"
    for row_idx, key in enumerate(["study", "group", "name", "unit", "abbr"], start=1):
        for col_idx, col in enumerate(merged_cols, start=1):
            ws_nice.cell(row=row_idx, column=col_idx, value=col[key] or None)
    for row_idx in range(data_row_count):
        for col_idx, col in enumerate(merged_cols, start=1):
            ws_nice.cell(row=row_idx + 6, column=col_idx, value=col["values"][row_idx])

    merge_repeated_cells(ws_nice, 1, 1, len(merged_cols))
    merge_header_row_with_fill(ws_nice, 2, 4, len(merged_cols))
    chol_start = next((idx for idx, col in enumerate(merged_cols, start=1) if col["abbr"] == "NLDLC"), None)
    if chol_start is not None:
        merge_header_row_with_fill(ws_nice, 3, chol_start, len(merged_cols) - 1)
    style_stratified_sheet(ws_nice, data_row_count + 5, len(merged_cols))
    wb_nice.save(OUT_NICE)

    wb_clean = Workbook()
    ws_clean = wb_clean.active
    ws_clean.title = "Clean"
    clean_headers = [col["abbr"] for col in merged_cols]
    ws_clean.append(clean_headers)
    for row_idx in range(data_row_count):
        ws_clean.append([col["values"][row_idx] for col in merged_cols])
    style_simple_sheet(ws_clean)
    wb_clean.save(OUT_CLEAN_XLSX)
    write_clean_csv(OUT_CLEAN_CSV, clean_headers, [[col["values"][row_idx] for col in merged_cols] for row_idx in range(data_row_count)])

    wb_codebook = Workbook()
    ws_code = wb_codebook.active
    ws_code.title = "Codebook"
    ws_code.append(
        [
            "abbreviated_variable_name",
            "variable_name",
            "variable_group_description",
            "measurement_units",
            "study_name",
        ]
    )
    for col in enrich_for_codebook(merged_cols):
        ws_code.append([col["abbr"], col["name"], col["group"], col["unit"], col["study"]])
    style_simple_sheet(ws_code)
    wb_codebook.save(OUT_CODEBOOK)

    print("Merged columns:", len(merged_cols))
    print("Rows:", data_row_count)
    print("Dropped duplicate variables:", ", ".join(duplicate_report) if duplicate_report else "None")
    print("Added variables:", ", ".join(col["abbr"] for col in add_cols))
    print("Outputs:")
    print(OUT_NICE.relative_to(PROJECT_ROOT))
    print(OUT_CLEAN_XLSX.relative_to(PROJECT_ROOT))
    print(OUT_CLEAN_CSV.relative_to(PROJECT_ROOT))
    print(OUT_CODEBOOK.relative_to(PROJECT_ROOT))


if __name__ == "__main__":
    main()
